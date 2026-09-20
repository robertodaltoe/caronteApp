"""
Logica dell'attività alternativa all'IRC (nota MIM prot. 11814 del
06/05/2026, punto 3.7): fabbisogno dalle ore di religione già presenti
in OrarioDocente, gruppi per slot, candidati in ordine di priorità.

Ordine di priorità della circolare, riprodotto in candidati():
  1. personale interamente o parzialmente a disposizione della scuola —
     qui: docenti con un'ora "a disposizione" in orario in quello slot, o
     che devono completare l'orario (ore in servizio < ore di contratto).
     I docenti di potenziamento NON rientrano (scelta di Roberto): la
     loro ora di potenziamento conta come occupata;
  2. in subordine, su base volontaria: supplenti con altro contratto a
     completamento, oppure docenti disponibili a ore eccedenti — solo chi
     ha dato la disponibilità (AlternativaIrcDisponibilita);
  3. in via residuale, un nuovo contratto ("da nominare").
In tutti i casi non è candidabile chi insegna in una delle classi del
gruppo (regola della circolare). Il docente è lo stesso per tutto l'anno,
quindi ogni verifica è sullo slot settimanale, non su date singole.
"""
import re
from collections import defaultdict
from datetime import date

from sqlalchemy import func

from models import db
from models.alternativa_irc import (
    AlternativaIrcAdesione, AlternativaIrcDisponibilita,
    AlternativaIrcGruppo, AlternativaIrcGruppoClasse,
)

# Soglia solo indicativa per segnalare un gruppo molto numeroso: non è un
# limite normativo, serve a far notare che forse conviene dividerlo.
SOGLIA_GRUPPO = 20

# Valore di Supplenza.classe per la copertura di un gruppo (che può
# raccogliere più classi): la lista delle classi sta in note/note_display.
CLASSE_SUPPLENZA = 'ALT. IRC'

GIORNI = ['Lunedì', 'Martedì', 'Mercoledì', 'Giovedì', 'Venerdì', 'Sabato']

LIVELLI = {
    1: 'Priorità 1 — a disposizione / completamento orario',
    2: 'Priorità 2 — disponibilità volontaria',
    3: 'Priorità 3 — nuovo contratto',
}


def norm_classe(c):
    return re.sub(r'\s+', '', str(c or '').strip().upper())


def label_classe(c):
    n = norm_classe(c)
    m = re.match(r'^(\d+[A-Z])(.+)$', n)
    return f'{m.group(1)} {m.group(2)}' if m else n


def anno_scol_di(d):
    return f'{d.year}-{d.year + 1}' if d.month >= 9 else f'{d.year - 1}-{d.year}'


# ── FABBISOGNO ────────────────────────────────────────────────────────

def slot_irc_per_classe():
    """{classe_compatta: {(giorno, ora), ...}} dalle ore di religione in orario."""
    from models.orario_docente import OrarioDocente
    out = defaultdict(set)
    righe = OrarioDocente.query.filter(
        func.upper(OrarioDocente.materia).like('%RELIG%')).all()
    for s in righe:
        if e_classe_reale(s.classe):
            out[norm_classe(s.classe)].add((s.giorno, s.ora))
    return out


def e_classe_reale(nome):
    """Esclude le pseudo-classi del potenziamento (es. '0A POT')."""
    n = (nome or '').strip().upper()
    return bool(n) and n[0].isdigit() and n[0] != '0' and 'POT' not in n


def tutte_le_classi():
    """Tutte le classi dell'istituto, non solo quelle con religione già in
    orario: con l'orario provvisorio l'ora di religione può mancare in
    alcune classi, che devono comunque poter ricevere l'adesione. Unione
    delle classi in orario e delle Assegnazioni (stessa fonte usata dal
    form degli eventi). Nessun numero fisso di classi o gruppi."""
    from models.orario_docente import OrarioDocente
    from models.assegnazione import AssegnazioneClasse
    classi = {norm_classe(c) for (c,) in
              OrarioDocente.query.with_entities(OrarioDocente.classe).distinct().all()
              if e_classe_reale(c)}
    classi |= {norm_classe(ac.label_classe) for ac in AssegnazioneClasse.query.all()
               if e_classe_reale(ac.label_classe)}
    return classi


def classi_con_adesione(anno):
    """Elenco di tutte le classi con l'adesione salvata. n_slot == 0 vuol
    dire che l'ora di religione non è (ancora) nell'orario importato."""
    slot = slot_irc_per_classe()
    ades = {a.classe: a for a in AlternativaIrcAdesione.query.filter_by(anno_scol=anno)}
    righe = []
    for cl in sorted(c for c in tutte_le_classi() | set(slot) | set(ades)
                     if e_classe_reale(c)):
        a = ades.get(cl)
        righe.append({
            'classe': cl, 'label': label_classe(cl),
            'n_slot': len(slot.get(cl, ())),
            'slot': sorted(slot.get(cl, ())),
            'n_con_docente': a.n_con_docente if a else 0,
            'n_altre': a.n_altre if a else 0,
            'note': a.note if a else '',
        })
    return righe


def salva_adesioni(anno, valori):
    """valori: {classe: (n_con_docente, n_altre, note)}"""
    for cl, (n_con, n_altre, note) in valori.items():
        a = AlternativaIrcAdesione.query.filter_by(anno_scol=anno, classe=cl).first()
        if not a:
            a = AlternativaIrcAdesione(anno_scol=anno, classe=cl)
            db.session.add(a)
        a.n_con_docente = max(0, n_con)
        a.n_altre = max(0, n_altre)
        a.note = (note or '').strip() or None
    db.session.commit()


def genera_gruppi(anno):
    """Crea/aggiorna un gruppo per ogni slot in cui almeno una classe con
    studenti da seguire ha religione. Non tocca mai il docente già
    assegnato: un gruppo che perde tutte le classi ma ha un docente resta
    (segnalato in 'senza_classi'), uno senza docente viene tolto."""
    slot = slot_irc_per_classe()
    ades = {a.classe: a for a in AlternativaIrcAdesione.query.filter_by(anno_scol=anno)
            if a.n_con_docente > 0}
    per_slot = defaultdict(set)
    classi_senza_slot = []
    for cl in ades:
        if not slot.get(cl):
            classi_senza_slot.append(cl)
        for s in slot.get(cl, ()):
            per_slot[s].add(cl)

    esistenti = {}
    for g in AlternativaIrcGruppo.query.filter_by(anno_scol=anno).order_by(AlternativaIrcGruppo.id):
        esistenti.setdefault((g.giorno, g.ora), g)

    creati = aggiornati = rimossi = 0
    senza_classi = []
    for s, classi in per_slot.items():
        g = esistenti.get(s)
        if not g:
            g = AlternativaIrcGruppo(anno_scol=anno, giorno=s[0], ora=s[1])
            db.session.add(g)
            db.session.flush()
            creati += 1
        else:
            aggiornati += 1
        attuali = {c.classe: c for c in g.classi}
        for cl in classi - set(attuali):
            db.session.add(AlternativaIrcGruppoClasse(id_gruppo=g.id, classe=cl))
        for cl, riga in attuali.items():
            if cl not in classi:
                db.session.delete(riga)

    for s, g in esistenti.items():
        if s in per_slot:
            continue
        if g.id_docente or g.da_nominare:
            for riga in list(g.classi):
                db.session.delete(riga)
            senza_classi.append(g)
        else:
            db.session.delete(g)
            rimossi += 1
    db.session.commit()
    return {'creati': creati, 'aggiornati': aggiornati, 'rimossi': rimossi,
            'senza_classi': len(senza_classi), 'classi_senza_slot': sorted(classi_senza_slot)}


def n_studenti_gruppo(gruppo, adesioni=None):
    if adesioni is None:
        adesioni = {a.classe: a for a in
                    AlternativaIrcAdesione.query.filter_by(anno_scol=gruppo.anno_scol)}
    return sum(adesioni[c.classe].n_con_docente for c in gruppo.classi if c.classe in adesioni)


# ── DISPONIBILITÀ DEI DOCENTI ─────────────────────────────────────────

class Contesto:
    """Dati precaricati una volta sola per valutare molti docenti/slot."""

    def __init__(self, anno):
        from models.orario_docente import OrarioDocente
        from models.orario_sostegno import OrarioSostegno
        from models.indisponibilita_ricorrente import IndisponibilitaRicorrente
        from routes.impostazione_anno import _docenti_per_anno

        self.anno = anno
        self.docenti = {d.id: d for d in _docenti_per_anno(anno)
                        if d.status_presenza not in ('aspettativa', 'ap_uscente')}
        self.occupati = defaultdict(set)      # id -> {(g, o)}
        self.disposizione = defaultdict(set)  # id -> {(g, o)}
        self.classi = defaultdict(set)        # id -> {classe compatta}
        self.n_ore = defaultdict(int)         # ore in servizio in orario
        for s in OrarioDocente.query.all():
            if s.tipo_ora == 'disposizione':
                self.disposizione[s.id_docente].add((s.giorno, s.ora))
                continue
            self.occupati[s.id_docente].add((s.giorno, s.ora))
            self.n_ore[s.id_docente] += 1
            if s.classe and s.classe[0].isdigit():
                self.classi[s.id_docente].add(norm_classe(s.classe))
        for s in OrarioSostegno.query.all():
            self.occupati[s.id_docente].add((s.giorno, s.ora))
            self.n_ore[s.id_docente] += 1
            if s.classe:
                self.classi[s.id_docente].add(norm_classe(s.classe))
        self.indisp = defaultdict(list)       # id -> [(giorno, ora|None)]
        for i in IndisponibilitaRicorrente.query.filter_by(attiva=True).all():
            self.indisp[i.id_docente].append((i.giorno, i.ora))
        self.gruppi = AlternativaIrcGruppo.query.filter_by(anno_scol=anno).all()

    def ore_alt(self, id_docente, escludi_gruppo=None):
        return sum(1 for g in self.gruppi
                   if g.id_docente == id_docente and g.id != escludi_gruppo)

    def slot_alt(self, id_docente, escludi_gruppo=None):
        return {(g.giorno, g.ora) for g in self.gruppi
                if g.id_docente == id_docente and g.id != escludi_gruppo}


def motivo_non_disponibile(ctx, id_docente, giorno, ora, classi, escludi_gruppo=None):
    """None se il docente può fare l'alternativa in quello slot, altrimenti
    la ragione (mostrata a Roberto)."""
    doc = ctx.docenti.get(id_docente)
    if not doc:
        return "non in servizio quest'anno"
    slot = (giorno, ora)
    if slot in ctx.occupati[id_docente]:
        return "ha già un impegno in orario in quell'ora (potenziamento incluso)"
    comuni = set(classi) & ctx.classi[id_docente]
    if comuni:
        return 'insegna nella classe ' + ', '.join(label_classe(c) for c in sorted(comuni))
    if doc.multi_sede:
        giorni = doc.giorni_presenza_list
        if giorni and giorno not in giorni:
            return f'{GIORNI[giorno]} è in servizio presso un\'altra scuola'
        uscita = doc.ora_uscita_map.get(giorno)
        if uscita is not None and ora > uscita:
            return 'in quell\'ora è già presso l\'altra sede'
    for g, o in ctx.indisp[id_docente]:
        if g == giorno and (o is None or o == ora):
            return 'indisponibilità ricorrente fissa'
    if slot in ctx.slot_alt(id_docente, escludi_gruppo):
        return 'già assegnato a un altro gruppo alla stessa ora'
    return None


def candidati(gruppo, ctx=None):
    """Candidati per un gruppo, già ordinati per priorità della circolare.
    Ritorna {'livelli': {1: [...], 2: [...]}, 'esclusi': [...]}; ogni voce
    è {'docente', 'dettaglio', 'ore_alt'}. Gli 'esclusi' sono solo i
    docenti che avevano dato disponibilità ma non vanno bene, con il
    motivo (per non far sembrare che siano stati dimenticati)."""
    ctx = ctx or Contesto(gruppo.anno_scol)
    classi = {c.classe for c in gruppo.classi}
    disp = {d.id_docente: d for d in
            AlternativaIrcDisponibilita.query.filter_by(anno_scol=gruppo.anno_scol)}
    livelli = {1: [], 2: []}
    esclusi = []
    for did, doc in ctx.docenti.items():
        motivo = motivo_non_disponibile(ctx, did, gruppo.giorno, gruppo.ora, classi, gruppo.id)
        d_flag = disp.get(did)
        if motivo:
            if d_flag:
                esclusi.append({'docente': doc, 'dettaglio': motivo})
            continue
        ore_alt = ctx.ore_alt(did, gruppo.id)
        voce = {'docente': doc, 'ore_alt': ore_alt}
        if (gruppo.giorno, gruppo.ora) in ctx.disposizione[did]:
            voce['dettaglio'] = 'a disposizione in quell\'ora'
            voce['ordine'] = 0
            livelli[1].append(voce)
            continue
        richieste = doc.ore_max_effettive_per_anno(gruppo.anno_scol)
        mancanti = richieste - (ctx.n_ore[did] + ore_alt)
        if ctx.n_ore[did] > 0 and mancanti > 0:
            voce['dettaglio'] = f'completa l\'orario (mancano {mancanti} h su {richieste})'
            voce['ordine'] = 1
            livelli[1].append(voce)
            continue
        if d_flag:
            if d_flag.max_ore is not None and ore_alt + 1 > d_flag.max_ore:
                esclusi.append({'docente': doc,
                                'dettaglio': f'raggiunto il massimo dichiarato ({d_flag.max_ore} h)'})
                continue
            nota = f', max {d_flag.max_ore} h' if d_flag.max_ore is not None else ''
            if d_flag.tipo == 'completamento':
                voce['dettaglio'] = 'supplente, completamento del contratto' + nota
                voce['ordine'] = 0
            else:
                voce['dettaglio'] = 'ore eccedenti' + nota
                voce['ordine'] = 1
            livelli[2].append(voce)
    for lst in livelli.values():
        lst.sort(key=lambda v: (v['ordine'], v['ore_alt'], v['docente'].cognome))
    return {'livelli': livelli, 'esclusi': esclusi}


def assegna(gruppo, id_docente=None, da_nominare=False, ctx=None):
    """Assegna il docente al gruppo (per tutto l'anno). Ritorna (ok, messaggio)."""
    if id_docente:
        cand = candidati(gruppo, ctx)
        validi = {v['docente'].id for lst in cand['livelli'].values() for v in lst}
        if id_docente not in validi:
            return False, 'Docente non candidabile per questo gruppo (vedi elenco).'
        gruppo.id_docente, gruppo.da_nominare = id_docente, False
    elif da_nominare:
        gruppo.id_docente, gruppo.da_nominare = None, True
    else:
        gruppo.id_docente, gruppo.da_nominare = None, False
    db.session.commit()
    return True, 'Assegnazione salvata.'


def avviso_docente_assegnato(gruppo, ctx):
    """Se il docente già assegnato non è più compatibile (es. l'orario è
    cambiato) ritorna il motivo, altrimenti None."""
    if not gruppo.id_docente:
        return None
    return motivo_non_disponibile(ctx, gruppo.id_docente, gruppo.giorno, gruppo.ora,
                                  {c.classe for c in gruppo.classi}, gruppo.id)


def gruppi_dettaglio(anno):
    ctx = Contesto(anno)
    ades = {a.classe: a for a in AlternativaIrcAdesione.query.filter_by(anno_scol=anno)}
    out = []
    for g in AlternativaIrcGruppo.query.filter_by(anno_scol=anno).order_by(
            AlternativaIrcGruppo.giorno, AlternativaIrcGruppo.ora).all():
        n = n_studenti_gruppo(g, ades)
        out.append({
            'gruppo': g,
            'classi': [label_classe(c) for c in g.classi_list],
            'n_studenti': n,
            'numeroso': n > SOGLIA_GRUPPO,
            'avviso': avviso_docente_assegnato(g, ctx),
            'cand': candidati(g, ctx),
        })
    return out


def riepilogo(anno):
    gruppi = AlternativaIrcGruppo.query.filter_by(anno_scol=anno).all()
    ades = AlternativaIrcAdesione.query.filter_by(anno_scol=anno).all()
    return {
        'n_classi_con_studenti': sum(1 for a in ades if a.n_con_docente > 0),
        'n_studenti': sum(a.n_con_docente for a in ades),
        'n_gruppi': len(gruppi),
        'n_assegnati': sum(1 for g in gruppi if g.id_docente),
        'n_da_nominare': sum(1 for g in gruppi if g.da_nominare),
        'n_scoperti': sum(1 for g in gruppi if not g.id_docente and not g.da_nominare),
        'n_disponibilita': AlternativaIrcDisponibilita.query.filter_by(anno_scol=anno).count(),
    }


# ── INTEGRAZIONE CON LE SUPPLENZE ─────────────────────────────────────

def docenti_occupati_alternativa(data_sel, ora):
    """Docenti impegnati nell'alternativa IRC in quel giorno/ora (l'incarico
    dura tutto l'anno scolastico): non vanno proposti per una supplenza."""
    gruppi = (AlternativaIrcGruppo.query
              .filter_by(anno_scol=anno_scol_di(data_sel), giorno=data_sel.weekday(), ora=ora)
              .filter(AlternativaIrcGruppo.id_docente.isnot(None)).all())
    return {g.id_docente for g in gruppi}


def gruppi_alternativa_docente(id_docente, data_sel, ore):
    """Gruppi con classi assegnati al docente in quel giorno, nelle ore date."""
    gruppi = (AlternativaIrcGruppo.query
              .filter_by(anno_scol=anno_scol_di(data_sel), giorno=data_sel.weekday(),
                         id_docente=id_docente)
              .filter(AlternativaIrcGruppo.ora.in_(list(ore))).all())
    return [g for g in gruppi if g.classi]


# ── EXPORT ────────────────────────────────────────────────────────────

def genera_xlsx(anno):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, color='FFFFFF')
    hdr_fill = PatternFill('solid', fgColor='7A1C17')

    def intesta(ws, colonne, larghezze):
        ws.append(colonne)
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill
        ws.freeze_panes = 'A2'
        for col, w in zip('ABCDEFGH', larghezze):
            ws.column_dimensions[col].width = w

    ws = wb.active
    ws.title = 'Gruppi'
    intesta(ws, ['Giorno', 'Ora', 'Classi', 'Studenti', 'Docente', 'Stato', 'Aula'],
            [12, 6, 38, 10, 26, 22, 14])
    for r in gruppi_dettaglio(anno):
        g = r['gruppo']
        if g.id_docente:
            doc = f'{g.docente.cognome} {g.docente.nome or ""}'.strip()
            stato = 'Assegnato'
        elif g.da_nominare:
            doc, stato = '', 'Supplente da nominare'
        else:
            doc, stato = '', 'Da assegnare'
        ws.append([GIORNI[g.giorno], g.ora, ', '.join(r['classi']), r['n_studenti'],
                   doc, stato, g.aula or ''])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(vertical='top', wrap_text=True)

    ws2 = wb.create_sheet('Adesioni')
    intesta(ws2, ['Classe', 'Studenti con docente', 'Altre scelte', 'Ore religione/sett.'],
            [12, 22, 14, 20])
    for r in classi_con_adesione(anno):
        if r['n_con_docente'] or r['n_altre']:
            ws2.append([r['label'], r['n_con_docente'], r['n_altre'], r['n_slot']])
    return wb
