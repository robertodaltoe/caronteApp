"""
modules/genera_docx_fse.py — Conversione dei documenti FSE/FESR
generati (vedi routes/progetti_fse.py) dall'HTML stampabile già usato
per il PDF (WeasyPrint) a un file .docx modificabile in Word, per chi
preferisce ricevere/archiviare il documento in quel formato invece del
PDF.

Non è un convertitore HTML->DOCX generico: riconosce solo il
sottoinsieme di markup usato nei template di
templates/progetti_fse/documenti/ (div con classi note, p, h2.articolo,
table, img, b, br) — è la stessa scelta già fatta per l'incorporazione
dei loghi come base64 in modules/dati_istituto.py, dove si preferisce
un meccanismo mirato e verificabile a una libreria generica di
conversione che non si può controllare riga per riga.
"""
import base64
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import html as lxml_html

PX = 0.75  # 1px CSS = 0.75pt: le dimensioni del PDF (WeasyPrint) si ricalcano tali e quali
LARGHEZZA_TESTO_CM = 17
INTERLINEA = 1.5

_ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT,
          'center': WD_ALIGN_PARAGRAPH.CENTER, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}

# Specchio di _stile_decreto.html: se cambia il CSS, aggiornare qui.
_STILE_CLASSE = {
    'protocollo': {'font-size': 10, 'text-align': 'left'},
    'oggetto': {'font-size': 10.5, 'font-weight': 700, 'margin': (14, 0, 4, 0), 'text-align': 'justify'},
    'estremi': {'font-size': 10.5, 'margin': (0, 0, 18, 0)},
    'qualifica': {'text-align': 'center', 'font-weight': 700, 'font-size': 11.5, 'margin': (20, 0, 10, 0)},
    'atto': {'text-align': 'center', 'font-weight': 700, 'font-size': 13, 'letter-spacing': 1.5,
             'margin': (22, 0, 14, 0)},
    'firma': {'margin': (34, 0, 0, 0), 'text-align': 'right', 'font-size': 10.5},
    'firma-nota': {'font-size': 8.5, 'color': '555555', 'text-align': 'right', 'margin': (2, 0, 0, 0)},
    'articolo': {'font-size': 11.5, 'font-weight': 700, 'margin': (18, 0, 6, 0)},
}
_BASE_P = {'text-align': 'justify', 'margin': (0, 0, 6, 0)}
_BASE_DIV = {'margin': (0, 0, 0, 0)}


def _decodifica_data_uri(src):
    if not src or not src.startswith('data:'):
        return None
    try:
        _, b64data = src.split(',', 1)
        return base64.b64decode(b64data)
    except (ValueError, base64.binascii.Error):
        return None


def _classi(el):
    return (el.get('class') or '').split()


def _num(v):
    m = re.match(r'\s*(-?[\d.]+)', str(v))
    return float(m.group(1)) if m else 0.0


def _parse_margin(v):
    """Ritorna (top, right, bottom, left) con valori in px o ('%', n)."""
    parti = v.split()
    def val(x):
        return ('%', _num(x)) if x.endswith('%') else _num(x)
    if len(parti) == 1:
        parti = parti * 4
    elif len(parti) == 2:
        parti = [parti[0], parti[1], parti[0], parti[1]]
    elif len(parti) == 3:
        parti = [parti[0], parti[1], parti[2], parti[1]]
    return tuple(val(x) for x in parti[:4])


def _proprieta(el, base=None):
    """Proprietà CSS effettive dell'elemento: base + classi note + style inline."""
    pr = dict(base or {})
    for c in _classi(el):
        pr.update(_STILE_CLASSE.get(c, {}))
    for dichiarazione in (el.get('style') or '').split(';'):
        if ':' not in dichiarazione:
            continue
        k, v = [x.strip() for x in dichiarazione.split(':', 1)]
        if k == 'margin':
            pr['margin'] = _parse_margin(v)
        elif k in ('margin-top', 'margin-bottom', 'margin-left'):
            m = list(pr.get('margin', (0, 0, 0, 0)))
            m[{'margin-top': 0, 'margin-bottom': 2, 'margin-left': 3}[k]] = _num(v)
            pr['margin'] = tuple(m)
        elif k == 'font-size':
            pr['font-size'] = _num(v)
        elif k == 'font-weight':
            pr['font-weight'] = 700 if v in ('bold', '700', '600', '800') else 400
        elif k == 'text-align':
            pr['text-align'] = v
        elif k == 'color':
            pr['color'] = v.lstrip('#')
    return pr


def _scrivi_run_inline(paragraph, el, pr=None):
    """Testo dell'elemento con <b>/<strong> in grassetto, <i>/<em> in
    corsivo, <br> come a capo, applicando dimensione/grassetto/colore
    dell'elemento contenitore."""
    pr = pr or {}
    def formatta(run, grassetto, corsivo):
        run.bold = True if (grassetto or pr.get('font-weight', 400) >= 700) else None
        run.italic = True if corsivo else None
        if pr.get('font-size'):
            run.font.size = Pt(pr['font-size'] * PX)
        col = pr.get('color')
        if col and re.fullmatch(r'[0-9a-fA-F]{6}', col):
            run.font.color.rgb = RGBColor.from_string(col.upper())
        if pr.get('letter-spacing'):
            rpr = run._r.get_or_add_rPr()
            sp = OxmlElement('w:spacing')
            sp.set(qn('w:val'), str(int(pr['letter-spacing'] * PX * 20)))
            rpr.append(sp)

    def aggiungi(testo, grassetto, corsivo):
        if not testo:
            return
        testo = re.sub(r'\s+', ' ', testo)
        formatta(paragraph.add_run(testo), grassetto, corsivo)

    def cammina(nodo, grassetto, corsivo):
        aggiungi(nodo.text, grassetto, corsivo)
        for figlio in nodo:
            if figlio.tag == 'br':
                paragraph.add_run().add_break()
            elif figlio.tag in ('b', 'strong'):
                cammina(figlio, True, corsivo)
            elif figlio.tag in ('i', 'em'):
                cammina(figlio, grassetto, True)
            else:
                cammina(figlio, grassetto, corsivo)
            aggiungi(figlio.tail, grassetto, corsivo)

    cammina(el, False, False)
    # rimuove lo spazio di troppo a inizio/fine paragrafo (a capo HTML)
    if paragraph.runs and paragraph.runs[0].text:
        paragraph.runs[0].text = paragraph.runs[0].text.lstrip(' ')
    for r in reversed(paragraph.runs):
        if r.text:
            r.text = r.text.rstrip(' ')
            break


def _formatta_paragrafo(paragraph, pr, larghezza_cm=LARGHEZZA_TESTO_CM):
    pf = paragraph.paragraph_format
    pf.line_spacing = INTERLINEA
    if pr.get('text-align') in _ALIGN:
        paragraph.alignment = _ALIGN[pr['text-align']]
    m = pr.get('margin', (0, 0, 0, 0))
    pf.space_before = Pt(_num(m[0]) * PX) if not isinstance(m[0], tuple) else None
    pf.space_after = Pt(_num(m[2]) * PX) if not isinstance(m[2], tuple) else None
    sx = m[3]
    if isinstance(sx, tuple):
        pf.left_indent = Cm(larghezza_cm * sx[1] / 100)
    elif sx:
        pf.left_indent = Pt(sx * PX)


def _sfondo_cella(cella, colore):
    tc_pr = cella._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), colore)
    tc_pr.append(shd)


def _bordi_tabella(tabella, colore='999999'):
    tbl_pr = tabella._tbl.tblPr
    bordi = OxmlElement('w:tblBorders')
    for lato in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{lato}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), colore)
        bordi.append(b)
    tbl_pr.append(bordi)
    marg = OxmlElement('w:tblCellMar')
    for lato, tw in (('top', 60), ('left', 90), ('bottom', 60), ('right', 90)):  # 4px / 6px
        e = OxmlElement(f'w:{lato}')
        e.set(qn('w:w'), str(tw))
        e.set(qn('w:type'), 'dxa')
        marg.append(e)
    tbl_pr.append(marg)


def _rendi_tabella(doc, el):
    righe = el.findall('tr')
    for sez in el.findall('tbody') + el.findall('thead'):
        righe += sez.findall('tr')
    if not righe:
        return
    classi = _classi(el)
    dimensione = 9.5 if 'tabella-incarichi' in classi else 9
    n_colonne = max(len([c for c in r if c.tag in ('td', 'th')]) for r in righe)
    tabella = doc.add_table(rows=0, cols=n_colonne)
    tabella.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordi_tabella(tabella)
    larghezze = [None] * n_colonne
    for riga_html in righe:
        celle_html = [c for c in riga_html if c.tag in ('td', 'th')]
        riga_docx = tabella.add_row()
        for i, cella_html in enumerate(celle_html):
            if i >= n_colonne:
                break
            cella = riga_docx.cells[i]
            par = cella.paragraphs[0]
            pr = {'font-size': dimensione, 'text-align': 'center' if 'num' in _classi(cella_html) else 'left',
                  'margin': (0, 0, 0, 0)}
            if cella_html.tag == 'th':
                pr['font-weight'] = 700
                _sfondo_cella(cella, 'EEEEEE')
            _scrivi_run_inline(par, cella_html, pr)
            _formatta_paragrafo(par, pr)
            w = re.search(r'width:\s*([\d.]+)px', cella_html.get('style') or '')
            if w and larghezze[i] is None:
                larghezze[i] = Pt(float(w.group(1)) * PX)
    if any(larghezze):
        tabella.autofit = False
        for riga in tabella.rows:
            for i, w in enumerate(larghezze):
                if w:
                    riga.cells[i].width = w
    spazio = doc.add_paragraph()
    spazio.paragraph_format.space_after = Pt(6)
    spazio.paragraph_format.space_before = Pt(0)


def _e_contenitore_di_blocchi(el):
    return any(figlio.tag in ('p', 'table', 'h2', 'div', 'ul', 'ol') for figlio in el)


def _rendi_lista(doc, el):
    for li in el.findall('li'):
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Pt(18 * PX + 12)
        par.paragraph_format.first_line_indent = Pt(-12)
        pr = {'text-align': 'justify', 'margin': (0, 0, 4, 0)}
        par.add_run('\u2022 ')
        _scrivi_run_inline(par, li, pr)
        _formatta_paragrafo(par, dict(pr, margin=(0, 0, 4, 0)))
        par.paragraph_format.left_indent = Pt(18 * PX + 12)


def html_a_docx(html_content, larghezza_intestazione_cm=17, larghezza_footer_cm=12):
    """Converte l'HTML stampabile di un documento FSE/FESR (vedi
    templates/progetti_fse/documenti/) in un .docx che ricalca la resa
    del PDF: stesse dimensioni carattere, interlinea, spaziature,
    allineamenti, corsivi/grassetti, rientri, tabelle con bordi e
    intestazioni ombreggiate, intestazione (solo prima pagina) e banner
    PN/UE (ogni pagina, footer nativo). Ritorna i byte del .docx."""
    albero = lxml_html.fromstring(html_content)
    corpo = albero.find('body')

    doc = Document()
    stile_normale = doc.styles['Normal']
    stile_normale.font.name = 'Times New Roman'
    stile_normale.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    stile_normale.font.size = Pt(11 * PX)
    stile_normale.paragraph_format.space_after = Pt(0)
    stile_normale.paragraph_format.space_before = Pt(0)
    stile_normale.paragraph_format.line_spacing = INTERLINEA

    sezione = doc.sections[0]
    sezione.page_width = Cm(21)
    sezione.page_height = Cm(29.7)
    sezione.left_margin = sezione.right_margin = Cm(2)
    sezione.top_margin = Cm(2.4)
    sezione.bottom_margin = Cm(4.6)
    sezione.header_distance = Cm(0.8)
    sezione.footer_distance = Cm(0.8)

    intestazione_bytes = None
    footer_bytes = None

    def cammina(el):
        nonlocal intestazione_bytes, footer_bytes
        classi = _classi(el)
        tag = el.tag

        if tag == 'div' and 'intestazione-istituto' in classi:
            img = el.find('.//img')
            if img is not None:
                intestazione_bytes = _decodifica_data_uri(img.get('src'))
            return
        if tag == 'div' and 'footer-ue' in classi:
            img = el.find('.//img')
            if img is not None:
                footer_bytes = _decodifica_data_uri(img.get('src'))
            return
        if tag in ('head', 'style', 'title'):
            return
        if tag == 'table':
            _rendi_tabella(doc, el)
            return
        if tag in ('ul', 'ol'):
            _rendi_lista(doc, el)
            return
        if tag == 'h2':
            paragrafo = doc.add_paragraph()
            pr = _proprieta(el, {'text-align': 'left', 'font-weight': 700, 'font-size': 11.5,
                                 'margin': (18, 0, 6, 0)})
            _scrivi_run_inline(paragrafo, el, pr)
            _formatta_paragrafo(paragrafo, pr)
            paragrafo.paragraph_format.keep_with_next = True
            return
        if tag == 'p':
            paragrafo = doc.add_paragraph()
            pr = _proprieta(el, _BASE_P)
            _scrivi_run_inline(paragrafo, el, pr)
            _formatta_paragrafo(paragrafo, pr)
            return
        if tag == 'div' and not _e_contenitore_di_blocchi(el):
            paragrafo = doc.add_paragraph()
            pr = _proprieta(el, _BASE_DIV)
            _scrivi_run_inline(paragrafo, el, pr)
            _formatta_paragrafo(paragrafo, pr)
            return
        for figlio in el:
            cammina(figlio)

    cammina(corpo)

    if intestazione_bytes or footer_bytes:
        sezione.different_first_page_header_footer = True
    if intestazione_bytes:
        intestazione = sezione.first_page_header
        paragrafo = intestazione.paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragrafo.paragraph_format.line_spacing = 1.0
        paragrafo.add_run().add_picture(io.BytesIO(intestazione_bytes), width=Cm(larghezza_intestazione_cm))
    if footer_bytes:
        for footer in ({sezione.footer, sezione.first_page_footer} if intestazione_bytes else {sezione.footer}):
            paragrafo = footer.paragraphs[0]
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragrafo.paragraph_format.line_spacing = 1.0
            paragrafo.add_run().add_picture(io.BytesIO(footer_bytes), width=Cm(larghezza_footer_cm))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
