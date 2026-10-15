"""
modules/piano_personale_xlsx.py — scheda Excel del Piano attività
personale, per i docenti che non possono usare il link (raggiungibile
solo dalla rete dell'istituto). Il docente compila SOLO la colonna
"Partecipo"; lo staff carica il file e Caronte riempie il piano.

Docente, anno e id degli eventi viaggiano nel file (foglio nascosto
'_meta' e colonna ID nascosta), così il file si riconosce da solo.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from models.attivita_ist import BUCKET_A, BUCKET_B, TIPI_ATTIVITA

FOGLIO_META = '_meta'
FOGLIO_PIANO = 'Piano'
RIGA_INTESTAZIONE = 8
SI = {'X', 'SI', 'SÌ', 'S', '1', 'TRUE', 'VERO'}

_FILL_HEAD = PatternFill('solid', fgColor='1F3864')
_FILL_INPUT = PatternFill('solid', fgColor='FFF2CC')


def nome_file_scheda(docente, anno_scol):
    base = ''.join(c for c in docente.cognome if c.isalnum()) or f'docente{docente.id}'
    return f'piano_attivita_{base}_{anno_scol}.xlsx'


def aggiungi_foglio_docente(wb, docente, anno_scol, eventi, selezionati, quota_a, quota_b,
                            primo=False):
    """Crea (o riusa il foglio attivo se primo=True) la scheda di un docente."""
    ws = wb.active if primo else wb.create_sheet()
    ws.title = FOGLIO_PIANO if primo else f'{docente.cognome[:20]}_{docente.id}'[:31]

    ws['A1'] = f'Piano attività personale a.s. {anno_scol}'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f'Docente: {docente.cognome} {docente.nome or ""}'.strip()
    ws['A2'].font = Font(bold=True)
    ws['A3'] = f'Quota ore bucket A: {quota_a}   |   Quota ore bucket B: {quota_b}'
    ws['A5'] = ('Scrivi X (o SI) nella colonna "Partecipo" per ogni impegno che sceglierai. '
                'Compila solo quella colonna: il resto è bloccato. Gli scrutini restano '
                'obbligatori e non sono in elenco.')
    ws['A5'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('A5:G6')

    intest = ['Partecipo', 'Data', 'Orario', 'Tipo', 'Descrizione', 'Bucket', 'Ore', 'ID']
    for i, t in enumerate(intest, 1):
        c = ws.cell(RIGA_INTESTAZIONE, i, t)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = _FILL_HEAD

    r = RIGA_INTESTAZIONE
    for e in eventi:
        r += 1
        orario = ' - '.join(x for x in (e.ora_inizio, e.ora_fine) if x)
        tipo = TIPI_ATTIVITA.get(e.tipo, {}).get('label', e.tipo)
        bucket = 'A' if e.bucket == BUCKET_A else 'B' if e.bucket == BUCKET_B else ''
        riga = [('X' if e.id in selezionati else ''), e.data, orario, tipo,
                e.titolo, bucket, round(e.durata_ore, 2), e.id]
        for i, v in enumerate(riga, 1):
            c = ws.cell(r, i, v)
            if i == 2:
                c.number_format = 'DD/MM/YYYY'
        inp = ws.cell(r, 1)
        inp.fill = _FILL_INPUT
        inp.protection = Protection(locked=False)

    if r > RIGA_INTESTAZIONE:
        dv = DataValidation(type='list', formula1='"X,SI"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'A{RIGA_INTESTAZIONE + 1}:A{r}')

    for col, w in zip('ABCDEFGH', (11, 12, 14, 24, 60, 8, 7, 8)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions['H'].hidden = True
    ws.freeze_panes = ws.cell(RIGA_INTESTAZIONE + 1, 1)
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    return ws


def _meta(wb, docenti_anno):
    meta = wb.create_sheet(FOGLIO_META)
    meta.append(['foglio', 'id_docente', 'anno_scol'])
    for nome_foglio, id_docente, anno in docenti_anno:
        meta.append([nome_foglio, id_docente, anno])
    meta.sheet_state = 'hidden'


def genera_xlsx(schede, anno_scol):
    """schede: lista di (docente, eventi, selezionati, quota_a, quota_b)."""
    wb = Workbook()
    meta = []
    for i, (d, eventi, sel, qa, qb) in enumerate(schede):
        ws = aggiungi_foglio_docente(wb, d, anno_scol, eventi, sel, qa, qb, primo=(i == 0))
        meta.append((ws.title, d.id, anno_scol))
    _meta(wb, meta)
    return wb


def leggi_xlsx(file_obj):
    """Ritorna una lista di dict, uno per foglio-docente riconosciuto:
    {id_docente, anno_scol, scelti:set(id_evento), id_presenti:set(id_evento)}.
    Solleva ValueError se il file non è una scheda di Caronte."""
    try:
        wb = load_workbook(file_obj, data_only=True)
    except Exception:
        raise ValueError('file non leggibile come Excel')
    if FOGLIO_META not in wb.sheetnames:
        raise ValueError('non è una scheda generata da Caronte (manca il foglio nascosto)')
    out = []
    for nome, id_docente, anno in list(wb[FOGLIO_META].iter_rows(min_row=2, values_only=True)):
        if not nome or nome not in wb.sheetnames or not id_docente:
            continue
        ws = wb[nome]
        scelti, presenti = set(), set()
        for r in range(RIGA_INTESTAZIONE + 1, ws.max_row + 1):
            eid = ws.cell(r, 8).value
            if not isinstance(eid, (int, float)):
                continue
            eid = int(eid)
            presenti.add(eid)
            v = str(ws.cell(r, 1).value or '').strip().upper()
            if v in SI:
                scelti.add(eid)
        out.append({'id_docente': int(id_docente), 'anno_scol': str(anno),
                    'scelti': scelti, 'id_presenti': presenti})
    if not out:
        raise ValueError('nessuna scheda docente trovata nel file')
    return out


NOME_CSV_CONDIVISIONE = 'elenco_condivisione.csv'

SCRIPT_CONDIVISIONE = """/**
 * Condivide ogni scheda Excel SOLO con il suo docente.
 * 1. Carica in una cartella Drive tutte le schede .xlsx e il file
 *    elenco_condivisione.csv (stessa cartella).
 * 2. Apri script.google.com con l'account della scuola, crea un progetto,
 *    incolla questo codice e sostituisci ID_CARTELLA (la parte finale
 *    dell'URL della cartella).
 * 3. Esegui condividiSchede() e autorizza. Il log mostra l'esito per riga.
 */
const ID_CARTELLA = 'ID_CARTELLA';

function condividiSchede() {
  const cartella = DriveApp.getFolderById(ID_CARTELLA);
  const csv = cartella.getFilesByName('elenco_condivisione.csv');
  if (!csv.hasNext()) throw new Error('elenco_condivisione.csv non trovato nella cartella');
  const righe = Utilities.parseCsv(csv.next().getBlob().getDataAsString());
  let ok = 0, saltate = 0;
  righe.slice(1).forEach(function (r) {
    const cognome = r[0], email = (r[2] || '').trim(), nomeFile = r[3];
    if (!email) { Logger.log('SALTATO (nessuna email): ' + cognome); saltate++; return; }
    const it = cartella.getFilesByName(nomeFile);
    if (!it.hasNext()) { Logger.log('SALTATO (file non trovato): ' + nomeFile); saltate++; return; }
    const f = it.next();
    f.setSharing(DriveApp.Access.PRIVATE, DriveApp.Permission.NONE);
    f.setShareableByEditors(false);
    f.addEditor(email);
    Logger.log('OK ' + cognome + ' -> ' + email);
    ok++;
  });
  Logger.log('Condivise: ' + ok + ' - Saltate: ' + saltate);
}
"""


def csv_condivisione(righe):
    """righe: lista di (cognome, nome, email, nome_file)."""
    import csv, io
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(['cognome', 'nome', 'email', 'file'])
    for r in righe:
        w.writerow(r)
    return buf.getvalue()
