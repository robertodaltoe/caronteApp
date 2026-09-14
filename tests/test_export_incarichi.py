"""
Test per l'export (PDF + XLSX) dell'elenco incarichi per docente
(Sessione 69 addendum 11): la pagina Report -> Report Docenti ->
Incarichi docenti non aveva nessuna funzione di stampa/esportazione.
"""
from models import db
from models.incarico import TipoIncarico, IncaricaDocente
from tests.conftest import crea_docente


def _crea_tabelle(app):
    with app.app_context():
        db.create_all()
    if 'report' not in app.blueprints:
        from routes.report import report_bp
        app.register_blueprint(report_bp)
    if not getattr(app, '_jinja_loader_reale', False):
        import os
        from jinja2 import FileSystemLoader
        templates_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'templates')
        app.jinja_loader = FileSystemLoader(templates_dir)
        app._jinja_loader_reale = True


def _crea_nomina(docente, anno='2026-2027', nome_incarico='Coordinatore di classe',
                  categoria='B', ore=None, importo=150.0, note=None,
                  anno_corso=3, sezione='A', indirizzo='RIM'):
    tipo = TipoIncarico(nome=nome_incarico, categoria=categoria, importo_default=importo)
    db.session.add(tipo)
    db.session.flush()
    n = IncaricaDocente(anno_scol=anno, id_tipo_incarico=tipo.id, id_docente=docente.id,
                         anno_corso=anno_corso, sezione=sezione, indirizzo=indirizzo,
                         importo=importo, note=note)
    db.session.add(n)
    db.session.commit()
    return n


def test_incarichi_pdf_con_dati(app, db_session):
    _crea_tabelle(app)
    with app.app_context():
        d = crea_docente('Santagata')
        _crea_nomina(d)

    with app.test_client() as c:
        r = c.get('/report/incarichi-docenti/pdf?anno=2026-2027')
        assert r.status_code == 200
        assert r.mimetype in ('application/pdf', 'text/html')
        if r.mimetype == 'text/html':
            assert 'Santagata' in r.get_data(as_text=True)


def test_incarichi_pdf_senza_dati_non_va_in_errore(app, db_session):
    _crea_tabelle(app)
    with app.test_client() as c:
        r = c.get('/report/incarichi-docenti/pdf?anno=2099-2100')
        assert r.status_code == 200


def test_incarichi_xlsx_con_dati(app, db_session):
    _crea_tabelle(app)
    with app.app_context():
        d = crea_docente('Santagata')
        _crea_nomina(d, note='Referente viaggi')

    with app.test_client() as c:
        r = c.get('/report/incarichi-docenti/xlsx?anno=2026-2027')
        assert r.status_code == 200
        assert r.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    import openpyxl, io
    wb = openpyxl.load_workbook(io.BytesIO(r.get_data()))
    ws = wb.active
    righe = [tuple(row) for row in ws.iter_rows(min_row=4, values_only=True) if row[0]]
    assert len(righe) == 1
    cognome, nome, incarico, categoria, contesto, ore, compenso, note = righe[0]
    assert cognome == 'Santagata'
    assert incarico == 'Coordinatore di classe'
    assert contesto == '3A RIM'
    assert note == 'Referente viaggi'


def test_incarichi_xlsx_senza_dati_solo_intestazione(app, db_session):
    _crea_tabelle(app)
    with app.test_client() as c:
        r = c.get('/report/incarichi-docenti/xlsx?anno=2099-2100')
        assert r.status_code == 200

    import openpyxl, io
    wb = openpyxl.load_workbook(io.BytesIO(r.get_data()))
    ws = wb.active
    righe = [row for row in ws.iter_rows(min_row=4, values_only=True) if row[0]]
    assert righe == []
