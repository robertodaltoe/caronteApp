from openpyxl import Workbook
from modules.parser_orario import parse_file


def test_orari_scritti_come_testo_sono_riconosciuti(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'ORARIO SETTIMANA 4b_teachers_ti'
    ws['C2'] = 'ORARIO PROVVISORIO'
    ws['E3'] = 'LUNEDì'
    ws['K3'] = 'MARTEDì'
    for col, t in zip('EFG', ['07:45 - 08:40', '08:40 - 09:35', '09:35 - 10:25']):
        ws[f'{col}4'] = t
    for col, t in zip('KLM', ['07:45 - 08:40', '08:40 - 09:35', '09:35 - 10:25']):
        ws[f'{col}4'] = t
    ws['C5'] = 'ROSSI'
    ws['E5'] = '2ALSC'; ws['E6'] = 'MATEMATICA'
    ws['F5'] = '---'
    ws['K5'] = '1ACAT'; ws['K6'] = 'MATEMATICA'
    p = tmp_path / 'orario.xlsx'
    wb.save(p)

    slots = parse_file(str(p))['slots']
    assert {(s['giorno'], s['ora'], s['classe']) for s in slots} == {
        (0, 1, '2ALSC'), (1, 1, '1ACAT')}
