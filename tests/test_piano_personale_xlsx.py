import io
from datetime import date
from types import SimpleNamespace as NS

import pytest
from openpyxl import load_workbook

from models.attivita_ist import BUCKET_A
from modules.piano_personale_xlsx import genera_xlsx, leggi_xlsx


def _ev(i):
    return NS(id=i, data=date(2026, 10, i), ora_inizio='14:30', ora_fine='16:30', tipo='collegio',
              titolo=f'Evento {i}', bucket=BUCKET_A, durata_ore=2.0)


def test_roundtrip_scelte():
    d = NS(id=7, cognome='ROSSI', nome='Anna')
    wb = genera_xlsx([(d, [_ev(1), _ev(2), _ev(3)], {2}, 20.0, 20.0)], '2026-2027')
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    wb2 = load_workbook(buf)
    ws = wb2['Piano']
    ws['A9'] = 'x'      # evento 1 scelto
    ws['A10'] = None    # evento 2 tolto
    buf2 = io.BytesIO(); wb2.save(buf2); buf2.seek(0)
    [s] = leggi_xlsx(buf2)
    assert (s['id_docente'], s['anno_scol']) == (7, '2026-2027')
    assert s['scelti'] == {1}
    assert s['id_presenti'] == {1, 2, 3}


def test_file_non_di_caronte_rifiutato():
    from openpyxl import Workbook
    buf = io.BytesIO(); Workbook().save(buf); buf.seek(0)
    with pytest.raises(ValueError):
        leggi_xlsx(buf)
